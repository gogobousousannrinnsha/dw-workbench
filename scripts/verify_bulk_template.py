"""Independent synthetic bulk-application and export probes.

No DocuWorks SDK, document rendering, OCR, or foreground GUI is exercised.
Run with --work-root outside the source checkout to retain diagnostic artifacts.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import csv
import ctypes
from dataclasses import asdict
import json
import os
from pathlib import Path
import sqlite3
import sys
import threading
import time
import traceback
import uuid
import zipfile

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "packages" / "dw-workbench"))
from dw_workbench.application import Workbench
from dw_workbench.domain import (Anchor, ExtractionProfile, FieldSchema, PageInfo,
    ResultSchema, RuleError, StaleRevision, Status, identifier, state_from)
from dw_workbench.exporting import export_dataset
from dw_workbench.storage import digest


def memory_bytes():
    """Current OS working set, rather than Python heap alone; optional on POSIX."""
    if os.name == "nt":
        class Counters(ctypes.Structure):
            _fields_ = [("cb", ctypes.c_ulong), ("PageFaultCount", ctypes.c_ulong),
                *[(name, ctypes.c_size_t) for name in ("PeakWorkingSetSize", "WorkingSetSize",
                    "QuotaPeakPagedPoolUsage", "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage",
                    "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage")]]
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        psapi = ctypes.WinDLL("psapi", use_last_error=True)
        kernel.GetCurrentProcess.restype = ctypes.c_void_p
        psapi.GetProcessMemoryInfo.argtypes = [ctypes.c_void_p, ctypes.POINTER(Counters), ctypes.c_ulong]
        psapi.GetProcessMemoryInfo.restype = ctypes.c_int
        values = Counters()
        values.cb = ctypes.sizeof(values)
        if psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(values), values.cb):
            return values.WorkingSetSize
        return None
    status = Path("/proc/self/status")
    if status.is_file():
        for line in status.read_text().splitlines():
            if line.startswith("VmRSS:"):
                return int(line.split()[1]) * 1024
    return None


@contextmanager
def measure(report, label):
    stop = threading.Event()
    samples = [memory_bytes()]
    def collect():
        while not stop.wait(0.025):
            samples.append(memory_bytes())
    sampler = threading.Thread(target=collect, daemon=True)
    before = time.perf_counter()
    sampler.start()
    try:
        yield
    finally:
        stop.set()
        sampler.join()
        samples.append(memory_bytes())
        valid = [value for value in samples if value is not None]
        report[label] = {"seconds": round(time.perf_counter() - before, 6),
            "working_set_before_bytes": samples[0], "working_set_after_bytes": samples[-1],
            "working_set_sampled_peak_bytes": max(valid) if valid else None,
            "sample_interval_seconds": 0.025}


def source(app, root, name, pages=3, geometries=None):
    path = root / name
    path.write_bytes(("Synthetic non-SDK document: " + name).encode("utf-8"))
    sid = app.prepare_source(path)
    app.finish_job("inspect-" + sid, {"pages": [asdict(page) for page in
        (geometries or [PageInfo(210, 297)] * pages)]})
    return sid


def template(app, *, count=3, scope="page", schema=None, profile=None, version=1):
    fields = tuple(FieldSchema("f" + str(index), "部品番号" if index == 0 else "項目" + str(index),
        "decimal" if index == 1 else "text", True, "℃" if index == 1 else "") for index in range(count))
    schema = schema or ResultSchema(identifier(), 1, "共通項目", fields)
    if not any(item.id == schema.id and item.version == schema.version for item in app.schemas()):
        app.register_schema(schema)
    p = ExtractionProfile(profile or identifier(), version, "合成テンプレート " + scope,
        (PageInfo(210, 297),), schema.fields,
        {f.id: {"page": 1, "rect": [10 + 35 * (index % 5), 20 + 15 * (index // 5), 30, 8]}
            for index, f in enumerate(schema.fields)}, scope, schema.id, schema.version)
    app.register_profile(p)
    return schema, p


def table_snapshot(app):
    names = ("records", "source_modes", "page_assignments", "jobs", "events", "candidates", "observations")
    return {name: app.store.rows(name) for name in names}


def complete(app, rid, marker="000125"):
    record, src, profile = app.context(rid)
    for field in profile.fields:
        record = app.store.record(rid)
        state = state_from(record["data"]["fields"][field.id])
        anchor = state.anchor or Anchor(src.id, src.sha256, record["page"] or 1, (10, 20, 30, 8))
        value = "179.8" if field.kind == "decimal" else marker
        app.edit(rid, field.id, record["revision"], value=value, raw=value, unit=state.unit, anchor=anchor)
        app.accept(rid, field.id, app.store.record(rid)["revision"])


def assignments(app, sid):
    return {a.page: a for a in app.assignments(sid)}


def single(app, sid, page, profile):
    mode = app.mode(sid)
    if mode["mode"] is None:
        app.set_mode(sid, profile.scope, mode["revision"])
    current = assignments(app, sid)[page]
    return app.assign_pages(sid, [page], profile.id, profile.version, {page: current.revision})[0]


def expect_failure(exception, operation):
    try:
        operation()
    except exception as error:
        return str(error)
    raise AssertionError("Expected " + exception.__name__)


def check(condition, description):
    if not condition:
        raise AssertionError(description)


def plan_entries(plan):
    return list(plan.entries)


def action_counts(plan):
    result = {}
    for entry in plan_entries(plan):
        result[entry.action] = result.get(entry.action, 0) + 1
    return result


def verify_excel_csv(root, dataset, results, expected_records):
    """Read XLSX through openpyxl and CSV directly, independently of writer helpers."""
    import openpyxl
    paths = {Path(row["path"]).suffix: root / row["path"] for row in results}
    check(all(row["status"] == "complete" for row in results), "Export returned a failed file")
    with zipfile.ZipFile(paths[".xlsx"]) as archive:
        check(archive.testzip() is None, "XLSX ZIP CRC")
        check(not any(b"<f>" in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/")), "Formula inserted")
    workbook = openpyxl.load_workbook(paths[".xlsx"], read_only=True, data_only=False)
    try:
        check(workbook.sheetnames == ["結果01", "根拠", "出力情報"], "Unexpected workbook sheets")
        rows = list(workbook["結果01"].iter_rows())
        check(len(rows) - 1 == expected_records, "Result row count")
        by_record = {cells[4].value: cells for cells in rows[1:]}
        expected = dataset["groups"][0]["records"]
        for record in expected:
            cells = by_record[record["id"]]
            check(cells[0].value == record["source"]["name"], "Source display association")
            check(cells[3].value == (record["page"] if record["page"] is not None else ""), "Page association")
            check(cells[8].value == record["fields"]["f0"]["value"] and cells[8].data_type == "s", "Leading-zero text type/value")
            check(cells[10].value == 179.8 and cells[10].data_type == "n", "Decimal numeric type/value")
        evidence = list(workbook["根拠"].iter_rows(values_only=True))
        header = evidence[0]
        records = {record["id"]: record for record in expected}
        def col(row, title):
            return row[header.index(title)]
        check(len(evidence) - 1 == sum(len(record["fields"]) for record in expected), "Evidence row count")
        for row in evidence[1:]:
            record = records[col(row, "記録ID")]
            state = record["fields"][col(row, "項目ID")]
            check(col(row, "原本ハッシュ") == record["source"]["sha256"], "Evidence source hash")
            check(col(row, "根拠ページ") == str(state["anchor"]["page"]), "Evidence actual page")
            check(col(row, "採用値") == state["value"], "Evidence adopted value")
    finally:
        workbook.close()
    with paths[".csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    check(paths[".csv"].read_bytes().startswith(b"\xef\xbb\xbf"), "CSV BOM")
    check(len(rows) - 1 == expected_records, "CSV row count")
    check(all(row[8].startswith("000") and row[10] == "179.8" for row in rows[1:]), "CSV exact decimal/text")
    return {"records": expected_records, "evidence_rows": len(evidence) - 1,
        "files": [{"name": Path(row["path"]).name, "sha256": digest(root / row["path"]),
                   "bytes": (root / row["path"]).stat().st_size} for row in results]}


def functional_probe(root):
    details = {}
    app = Workbench(root / "案件 一括適用")
    try:
        schema, page_profile = template(app)
        _, version_two = template(app, schema=schema, profile=page_profile.id, version=2)
        _, document_profile = template(app, schema=schema, scope="document")
        a = source(app, root, "帳票 A.xdw", pages=3)
        b = source(app, root, "帳票 B.xdw", pages=2)
        before = table_snapshot(app)
        plan = app.plan_template_application(page_profile.id, 1, {a: [1, 2, 3], b: [1, 2]})
        check(table_snapshot(app) == before, "Preview changed database")
        check(action_counts(plan) == {"apply": 5}, "Expected five applicable page records")
        app.apply_template_plan(plan)
        check(len(app.store.rows("records")) == 5, "Five records were not created")
        check(all(app.mode(sid)["mode"] == "page" for sid in (a, b)), "Page units initialization")
        check(all(app.store.record(row["id"])["data"]["fields"]["f0"]["anchor"]["page"] == row["page"] for row in app.store.rows("records")), "Template relative anchor mapping")
        check(not any(row["kind"] in ("render", "ocr") for row in app.store.rows("jobs")), "Application automatically started OCR/rendering")
        details["multi_document_preview_apply"] = {"records": 5, "preview_read_only": True, "automatic_ocr_jobs": 0}

        old = assignments(app, a)[1].current_record_id
        complete(app, old)
        old_data = app.store.record(old)["data"]
        app.exclude_pages(a, [2], "説明ページ", {2: assignments(app, a)[2].revision})
        before = table_snapshot(app)
        protected = app.plan_template_application(version_two.id, 2, {a: [1, 2, 3], b: [1, 2]})
        check(action_counts(protected) == {"skip": 5}, "Default must protect all existing/excluded assignments")
        app.apply_template_plan(protected)
        check(table_snapshot(app) == before, "Skipped application changed records")
        details["unassigned_only_protection"] = {"skipped": 5, "accepted_and_excluded_unchanged": True}

        jobs = app.ocr_jobs(old)
        app.start_job(jobs[0])
        inclusive = app.plan_template_application(version_two.id, 2, {a: [1]}, unassigned_only=False)
        app.apply_template_plan(inclusive)
        changed = assignments(app, a)[1].current_record_id
        check(changed != old and app.store.record(old)["active"] == 0, "Version switch did not archive old record")
        check(app.store.record(old)["data"] == old_data, "Archived confirmation data changed")
        check(all(state["value"] == "" and state["status"] == Status.MISSING for state in app.store.record(changed)["data"]["fields"].values()), "New version inherited adopted values")
        check(all(app.job(job)["status"] == "cancelled" for job in jobs), "Old OCR jobs were not cancelled")
        app.finish_job(jobs[0], {"text": "late", "regions": [], "engine": {}})
        check(not app.candidates(changed, "f0"), "Late old OCR contaminated new record")
        same = app.plan_template_application(version_two.id, 2, {a: [1]}, unassigned_only=False)
        before = table_snapshot(app)
        check(action_counts(same) == {"same": 1}, "Same version not classified as unchanged")
        app.apply_template_plan(same)
        check(table_snapshot(app) == before, "Same version changed database")
        details["version_switch_and_same_version"] = {"history_preserved": True, "old_jobs_cancelled": len(jobs), "same_version_noop": True}

        stale = app.plan_template_application(page_profile.id, 1, {a: [1], b: [1]}, unassigned_only=False)
        r = app.store.record(changed)
        app.edit(changed, "f0", r["revision"], value="edited-after-preview", raw="edited-after-preview", unit="", anchor=state_from(r["data"]["fields"]["f0"]).anchor)
        before = table_snapshot(app)
        message = expect_failure(StaleRevision, lambda: app.apply_template_plan(stale))
        check(table_snapshot(app) == before, "Stale record preview wrote other targets")
        stale_assignment = app.plan_template_application(page_profile.id, 1, {a: [2], b: [2]}, unassigned_only=False)
        app.exclude_pages(a, [2], "理由を変更", {2: assignments(app, a)[2].revision})
        before = table_snapshot(app)
        expect_failure(StaleRevision, lambda: app.apply_template_plan(stale_assignment))
        check(table_snapshot(app) == before, "Stale assignment preview wrote other targets")
        c = source(app, root, "方式変更 C.xdw", pages=2)
        stale_mode = app.plan_template_application(page_profile.id, 1, {c: [1]})
        app.set_mode(c, "document")
        before = table_snapshot(app)
        expect_failure(StaleRevision, lambda: app.apply_template_plan(stale_mode))
        check(table_snapshot(app) == before, "Stale processing-unit preview wrote targets")
        details["stale_preview"] = {"record_revision": True, "assignment_revision": True, "mode_revision": True, "diagnostic": message}

        locked = source(app, root, "文書全体 D.xdw", pages=1)
        doc_record = single(app, locked, 0, document_profile)
        complete(app, doc_record)
        mixed = app.plan_template_application(page_profile.id, 1, {locked: [1], c: [1, 2], a: [99]})
        check(action_counts(mixed) == {"skip": 2, "apply": 2}, "Locked scope or out-of-range target not skipped")
        app.apply_template_plan(mixed)
        check(app.mode(locked)["mode"] == "document" and app.store.record(doc_record)["active"] == 1, "Locked document unit changed")
        check(app.mode(c)["mode"] == "page", "Unlocked document unit was not converted")
        details["mixed_processing_units_and_invalid_page"] = {"locked_document_preserved": True, "unlocked_unit_changed": True, "invalid_page_skipped": True}

        mismatch = source(app, root, "寸法不一致 E.xdw", geometries=[PageInfo(297, 420)])
        plan = app.plan_template_application(page_profile.id, 1, {mismatch: [1]})
        app.apply_template_plan(plan)
        mrid = assignments(app, mismatch)[1].current_record_id
        check(not app.store.record(mrid)["data"]["geometry_matches"], "Geometry mismatch not retained")
        check(all(value["anchor"] is None for value in app.store.record(mrid)["data"]["fields"].values()), "Mismatch inherited incorrect anchors")
        expect_failure(RuleError, lambda: app.ocr_jobs(mrid))
        complete(app, mrid, marker="000126")
        details["geometry_mismatch_manual_completion"] = {"fields_retained": 3, "automatic_ocr_rejected": True, "manual_completion": True}

        rollback = source(app, root, "強制失敗 F.xdw", pages=2)
        earlier = assignments(app, b)[1].current_record_id
        complete(app, earlier)
        running = app.ocr_jobs(earlier)[0]
        app.start_job(running)
        plan = app.plan_template_application(version_two.id, 2, {b: [1], rollback: [1, 2]}, unassigned_only=False)
        before = table_snapshot(app)
        check("'" not in rollback, "Unexpected fixture ID")
        app.store.db.execute("CREATE TEMP TRIGGER synthetic_failure BEFORE INSERT ON records WHEN NEW.source_id='" + rollback + "' BEGIN SELECT RAISE(ABORT,'synthetic second document failure'); END")
        try:
            error = expect_failure(sqlite3.IntegrityError, lambda: app.apply_template_plan(plan))
        finally:
            app.store.db.execute("DROP TRIGGER synthetic_failure")
        check(table_snapshot(app) == before, "Failure did not roll back modes, records, jobs, assignments and events")
        check(app.job(running)["status"] == "running", "Rollback left OCR cancelled")
        details["whole_batch_rollback"] = {"tables_unchanged": sorted(before), "diagnostic": error}

        # Export complete representatives while clearly recording unfinished pages.
        complete(app, changed)
        frozen = app.dataset(app.finalize(completed_only=True))
        check(len(frozen["groups"]) == 1, "Common schema did not group layouts/page/document units")
        count = len(frozen["groups"][0]["records"])
        results = export_dataset(app.store, frozen, ("xlsx", "csv"))
        details["case_output"] = verify_excel_csv(app.store.root, frozen, results, count)
        details["case_output"]["excluded_records"] = len(frozen["excluded"])
        check(app.store.db.execute("PRAGMA integrity_check").fetchone()[0] == "ok", "SQLite integrity")
        check(app.store.db.execute("PRAGMA foreign_key_check").fetchall() == [], "SQLite foreign keys")
        project_root = app.store.root
        app.close()
        app = Workbench(project_root)
        check(app.dataset(frozen["id"]) == frozen, "Frozen dataset changed after reopening")
        check(app.store.record(old)["data"] == old_data, "Archived accepted values changed after reopening")
        details["save_reopen"] = {"frozen_dataset_unchanged": True, "archived_data_unchanged": True}
    finally:
        app.close()
    return details


def load_probe(root):
    details = {"documents": 100, "pages_per_document": 10, "fields_per_page": 30,
        "fixture_confirmation": "domain state acceptance + one Store.update transaction; no UI/OCR benchmark"}
    app = Workbench(root / "負荷案件 100文書")
    try:
        _, profile = template(app, count=30)
        with measure(details, "register_sources"):
            ids = [source(app, root, "負荷 " + str(index).zfill(3) + ".xdw", pages=10) for index in range(100)]
        targets = {sid: list(range(1, 11)) for sid in ids}
        with measure(details, "preview"):
            plan = app.plan_template_application(profile.id, 1, targets)
        check(action_counts(plan) == {"apply": 1000}, "Load preview row count")
        with measure(details, "apply_and_save"):
            app.apply_template_plan(plan)
        records = app.store.rows("records")
        check(len(records) == 1000 and sum(len(app.store.record(r["id"])["data"]["fields"]) for r in records) == 30000, "Load fields count")
        with measure(details, "fixture_confirm_all_fields"):
            with app.store.transaction():
                for r in records:
                    record, src, p = app.context(r["id"])
                    fields = {}
                    for f in p.fields:
                        state = state_from(record["data"]["fields"][f.id])
                        state.value = state.raw = "179.8" if f.kind == "decimal" else "000125"
                        state.accept(f, src, record["revision"] + 1)
                        fields[f.id] = asdict(state)
                    app.store.update(record["id"], record["revision"], record["data"] | {"fields": fields}, "synthetic-confirmation-fixture")
        with measure(details, "finalize"):
            dataset = app.dataset(app.finalize())
        with measure(details, "xlsx_csv_generate_and_application_verify"):
            results = export_dataset(app.store, dataset, ("xlsx", "csv"))
        with measure(details, "independent_output_read"):
            details["output"] = verify_excel_csv(app.store.root, dataset, results, 1000)
        project_root = app.store.root
        app.close()
        with measure(details, "reopen"):
            app = Workbench(project_root)
            check(len(app.store.rows("records")) == 1000, "Load reopen record count")
            check(len(app.dataset(dataset["id"])["groups"][0]["records"]) == 1000, "Load reopen frozen records")
        details["database_bytes"] = (project_root / "project.sqlite").stat().st_size
        check(app.store.db.execute("PRAGMA integrity_check").fetchone()[0] == "ok", "Load SQLite integrity")
        check(app.store.db.execute("PRAGMA foreign_key_check").fetchall() == [], "Load SQLite foreign keys")
    finally:
        app.close()
    return details


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work-root", type=Path, default=Path(os.environ.get("DW_WORKBENCH_VERIFY_ROOT", "bulk-template-verification")))
    parser.add_argument("--skip-load", action="store_true")
    args = parser.parse_args()
    root = args.work_root.resolve() / ("run-" + uuid.uuid4().hex[:12])
    root.mkdir(parents=True, exist_ok=False)
    report = {"scope": "synthetic non-SDK XDW bytes; application/storage/export only", "python": sys.version,
        "work_root": str(root), "source_revision": os.environ.get("DW_WORKBENCH_SOURCE_REVISION", "working-tree")}
    try:
        report["functional"] = functional_probe(root)
        if not args.skip_load:
            report["load"] = load_probe(root)
        report["status"] = "passed"
    except BaseException as error:
        report["status"] = "failed"
        report["error"] = str(error)
        report["traceback"] = traceback.format_exc()
    path = root / "verification.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": report["status"], "report": str(path), "error": report.get("error")}, ensure_ascii=False))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
