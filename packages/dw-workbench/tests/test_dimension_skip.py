"""Derived dimension skips retain assignments, human evidence and frozen output."""
from dataclasses import replace
import pytest
from conftest import complete
from test_v020_application import source, templates, assign
from dw_workbench.application import Workbench
from dw_workbench.domain import Anchor, DimensionMismatch, PageInfo, RuleError, identifier, state_from
from dw_workbench.exporting import info_rows_v2, write_xlsx_v2
from dw_workbench.storage import digest, encode


def mixed(app, workdir):
    sid = source(app, workdir, pages=[PageInfo(210, 297), PageInfo(297, 420), PageInfo(210, 297)])
    _, profiles = templates(app)
    app.set_mode(sid, 'page')
    return sid, profiles[0], assign(app, sid, [1, 2, 3], profiles[0])


def complete_manual_mismatch(app, rid):
    record, original, _ = app.context(rid)
    for field, value in (('part', '001234'), ('temp', '25')):
        record = app.store.record(rid)
        app.edit(rid, field, record['revision'], value=value, raw=value,
            unit=record['data']['fields'][field]['unit'],
            anchor=Anchor(original.id, original.sha256, record['page'] or 1, (10, 20, 30, 8)))
        app.accept(rid, field, app.store.record(rid)['revision'])
    from dw_workbench.domain import Status
    app.mark(rid, 'note', app.store.record(rid)['revision'], Status.NOT_APPLICABLE, '記載なし')


def test_mixed_navigation_skips_only_mismatches_and_keeps_unassigned(app, workdir):
    sid, profile, (first, skipped, third) = mixed(app, workdir)
    other = source(app, workdir, name='unassigned.xdw', pages=[PageInfo(297, 420)])
    app.set_mode(other, 'page')
    before = app.store.record(skipped)
    assert {t['record_id'] for t in app.review_targets() if t['record_id']} == {first, third}
    assert app.next_unfinished((sid, 2, 'part'))['record_id'] == third
    assert app.next_unfinished((sid, 2, None))['record_id'] == third
    complete(app, first)
    complete(app, third)
    target = app.next_unfinished((sid, 3, 'note'))
    assert target['source_id'] == other and target['record_id'] is None
    assert app.store.record(skipped) == before
    assert app.assignments(sid)[1].state == 'applied'
    assert app.workflow_skips()[0]['record_id'] == skipped
    assert next(x for x in app.incomplete() if x['record_id'] == skipped)['kind'] == 'dimension_mismatch'


@pytest.mark.parametrize('page,matching', [
    (PageInfo(210.009, 297), True), (PageInfo(210.011, 297), False),
    (PageInfo(210, 297.009), True), (PageInfo(210, 297.011), False),
    (PageInfo(210, 297, 90), False),
])
def test_original_geometry_tolerance_and_rotation_are_preserved(app, workdir, page, matching):
    sid = source(app, workdir, pages=[page])
    _, profiles = templates(app)
    app.set_mode(sid, 'page')
    rid = assign(app, sid, [1], profiles[0])[0]
    assert profiles[0].matches(app.source(sid), 1) is matching
    assert (app.workflow_skip(rid) is None) is matching


def test_document_mismatch_skips_whole_unit_without_empty_jobs_or_loop(app, workdir):
    sid = source(app, workdir, pages=[PageInfo(210, 297), PageInfo(210, 297)])
    _, profiles = templates(app)
    profile = replace(profiles[0], id=identifier(), scope='document')
    app.register_profile(profile)  # one-page template against a two-page document
    rid = app.apply_profile(sid, profile.id, 1)
    before = app.store.record(rid)
    for _ in range(3):
        assert app.next_unfinished() is None
        with pytest.raises(DimensionMismatch):
            app.ocr_jobs(rid)
    assert not [j for j in app.store.rows('jobs') if j['kind'] == 'ocr']
    assert app.store.record(rid) == before
    assert len(app.workflow_skips()) == 1


def test_frozen_exclusion_reason_manual_values_and_reassignment_survive_restart(app, workdir):
    sid, profile, (first, skipped, third) = mixed(app, workdir)
    complete(app, first)
    complete(app, third)
    with pytest.raises(RuleError):
        app.finalize()  # existing explicit completed-only policy remains
    frozen_id = app.finalize(completed_only=True)
    frozen = app.dataset(frozen_id)
    assert [r['id'] for r in frozen['groups'][0]['records']] == [first, third]
    exclusion = frozen['excluded'][0]
    assert exclusion['record_id'] == skipped and exclusion['kind'] == 'dimension_mismatch'
    assert '寸法' in exclusion['reason'] and profile.name in exclusion['reason']
    rows = [[value for _, value in row] for row in info_rows_v2(frozen, 'synthetic')]
    assert any('寸法不一致スキップ' in str(row) for row in rows)
    xlsx = workdir/'synthetic-exclusions.xlsx'
    write_xlsx_v2(xlsx, frozen, workdir)
    from openpyxl import load_workbook
    book = load_workbook(xlsx, read_only=True)
    try:
        assert any('寸法不一致スキップ' in str(row) for row in book['出力情報'].values)
        assert book['結果01'].max_row == 3
    finally:
        book.close()
    complete_manual_mismatch(app, skipped)
    manual = app.store.record(skipped)
    manual_dataset_id = app.finalize()
    manual_dataset = app.dataset(manual_dataset_id)
    assert len(manual_dataset['groups'][0]['records']) == 3
    original_hash = digest(app.store.path(app.source(sid).path))
    fit = replace(profile, id=identifier(), name='適合A3', pages=(PageInfo(297, 420),))
    app.register_profile(fit)
    newer = assign(app, sid, [2], fit)[0]
    assert app.workflow_skip(newer) is None
    assert app.next_unfinished()['record_id'] == newer
    assert app.store.record(skipped) == manual | {'active': 0}
    assert state_from(app.store.record(newer)['data']['fields']['part']).anchor.page == 2
    assert app.dataset(frozen_id) == frozen and app.dataset(manual_dataset_id) == manual_dataset
    assert digest(app.store.path(app.source(sid).path)) == original_hash
    root = app.store.root
    app.close()
    reopened = Workbench(root)
    try:
        assert not reopened.workflow_skips()
        assert reopened.next_unfinished()['record_id'] == newer
        assert reopened.dataset(frozen_id) == frozen
        assert reopened.dataset(manual_dataset_id) == manual_dataset
        assert reopened.store.record(skipped)['data'] == manual['data']
    finally:
        reopened.close()


@pytest.mark.parametrize('action', ['pending', 'start', 'finish'])
def test_legacy_geometry_flags_and_queued_results_cannot_bypass_skip(app, workdir, action):
    sid, profile, (_, skipped, _) = mixed(app, workdir)
    before = app.store.record(skipped)
    # Legacy flag is advisory; saved source/template geometry remains authoritative.
    data = dict(before['data'])
    data.pop('geometry_matches')
    app.store.db.execute('UPDATE records SET data=? WHERE id=?', (encode(data), skipped))
    assert app.workflow_skip(skipped)
    record = app.store.record(skipped)
    job_data = dict(record_id=skipped, source_id=sid, sha256=app.source(sid).sha256,
        profile_id=profile.id, profile_version=1, page=2, field_id='part',
        assignment_revision=record['assignment_revision'], rect=[10, 20, 30, 8])
    job = app.add_job('ocr', job_data)
    if action == 'pending':
        assert job not in [j['id'] for j in app.pending_jobs()]
    elif action == 'start':
        assert app.start_job(job) is None
    else:
        app.store.db.execute("UPDATE jobs SET status='running' WHERE id=?", (job,))
        app.finish_job(job, {'text': 'LATE', 'regions': [], 'engine': {}})
    assert app.job(job)['status'] == 'cancelled'
    assert 'スキップ' in app.job(job)['error']
    assert not app.candidates(skipped, 'part')
    assert app.store.record(skipped)['data'] == data
    assert assign(app, sid, [2], profile) == [skipped]
    root = app.store.root
    app.close()
    reopened = Workbench(root)
    try:
        assert reopened.workflow_skip(skipped)
        assert reopened.next_unfinished((sid, 2, 'part'))['page'] == 3
    finally:
        reopened.close()
