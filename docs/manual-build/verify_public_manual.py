"""Read-only checks for the public manual and its synthetic examples.

Run in this build folder or pass --manual to a copied consumer folder. Visual
PDF review is a separate step; this script never launches the application.
"""
import argparse
import csv
import hashlib
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook
from pypdf import PdfReader


PRIVATE = re.compile(r"(?<![A-Za-z0-9])[A-Za-z]:[/\\]|/(?:Users|home)/")


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids, self.references = set(), []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if attrs.get("id"):
            self.ids.add(attrs["id"])
        for key in ("href", "src"):
            if attrs.get(key):
                self.references.append(attrs[key])


def main():
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--manual", type=Path, default=Path(__file__).resolve().parent)
    opts = args.parse_args()
    root = opts.manual.resolve()
    records = root / "検証記録" if (root / "検証記録").is_dir() else root
    demo = json.loads((records / "beginner_demo.json").read_text(encoding="utf-8-sig"))
    checks = []

    def check(name, condition, detail=None):
        checks.append({"name": name, "passed": bool(condition), "detail": detail})

    pdf = root / "DW-Workbench_v0.2.0_図解操作マニュアル.pdf"
    reader = PdfReader(pdf)
    check("PDF pages", len(reader.pages) == 128, len(reader.pages))
    text = "\n".join(page.extract_text() for page in reader.pages)
    check("PDF no private absolute paths", not PRIVATE.search(text))
    parsed = Links()
    html = (root / "index.html").read_text(encoding="utf-8-sig")
    parsed.feed(html)
    check("HTML no private absolute paths", not PRIVATE.search(html))
    control_ids = {n for n in parsed.ids if re.fullmatch(r"(?:common|viewer|worklist|setup|review|history|library)_\d+", n)}
    check("all 116 control anchors", len(control_ids) == 116, len(control_ids))
    bad, files = [], set()
    for raw in parsed.references:
        url = urlsplit(raw)
        if url.scheme or url.netloc:
            continue
        if url.path:
            p = (root / unquote(url.path)).resolve()
            if not p.is_relative_to(root) or not p.is_file():
                bad.append(raw)
            else:
                files.add(p.relative_to(root).as_posix())
        if url.fragment and not url.path and unquote(url.fragment) not in parsed.ids:
            bad.append(raw)
    check("offline local references", not bad, {"file_count": len(files), "invalid": bad})
    source = root / demo["input"]
    check("synthetic source unchanged", sha(source) == demo["original_sha256"])
    xlsx, csv_path = root / demo["xlsx"], root / demo["csv"]
    book = load_workbook(xlsx, read_only=True, data_only=False)
    sheet = book["結果01"]
    check("part text with leading zeros", sheet["I2"].value == "000125" and sheet["I2"].data_type == "s")
    check("temperature numeric", sheet["K2"].value == 179.8 and sheet["K2"].data_type == "n")
    check("unit separate", sheet["J2"].value in (None, "") and sheet["L2"].value == "℃")
    check("one source one page one row", sheet.max_row == 2 and sheet["A2"].value == "入門サンプル.xdw" and sheet["D2"].value == 1)
    check("evidence and output sheets", book.sheetnames == ["結果01", "根拠", "出力情報"])
    book.close()
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    check("CSV exact decimals and zeros", len(rows) == 1 and rows[0]["部品番号"] == "000125" and rows[0]["測定温度"] == "179.8" and rows[0]["測定温度 単位"] == "℃")
    check("CSV UTF-8 BOM", csv_path.read_bytes().startswith(b"\xef\xbb\xbf"))
    private = []
    for p in root.rglob("*"):
        if p.is_file() and p.suffix in (".json", ".py", ".md", ".txt", ".html"):
            if PRIVATE.search(p.read_text(encoding="utf-8-sig")):
                private.append(p.relative_to(root).as_posix())
    check("public text privacy scan", not private, private)
    sensitive_binary = []
    tokens = (":\\Users", ":/Users", "codex-runtimes")
    for p in (source, xlsx):
        contents = [(p.name, p.read_bytes())]
        if p.suffix == ".xlsx":
            with zipfile.ZipFile(p) as z:
                contents += [(name, z.read(name)) for name in z.namelist()]
        for name, content in contents:
            if any(t.encode(enc) in content for t in tokens for enc in ("utf-8", "utf-16le")):
                sensitive_binary.append(name)
    check("synthetic binary privacy scan", not sensitive_binary, sensitive_binary)
    report = {"scope": "Read-only structural, offline-link, type and privacy checks; PDF visual QA is separate.",
              "pdf": pdf.name, "pdf_sha256": sha(pdf), "checks": checks,
              "all_checks_passed": all(c["passed"] for c in checks),
              "sample_hashes": {p.relative_to(root).as_posix(): sha(p) for p in (source, xlsx, csv_path)}}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["all_checks_passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
